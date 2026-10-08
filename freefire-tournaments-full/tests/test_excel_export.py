import io
from datetime import datetime, timedelta
from decimal import Decimal
import openpyxl
from werkzeug.security import generate_password_hash
from app import db
from app.models.admin import Admin
from app.models.tournament import Tournament
from app.models.registration import Registration
from app.models.player import Player


def _setup_data(app):
    with app.app_context():
        admin = Admin(
            email="admin@test.com",
            password_hash=generate_password_hash("password123"),
        )
        t = Tournament(
            name="Excel Cup",
            description="desc",
            entry_fee=Decimal("50.00"),
            prize_pool=Decimal("5000.00"),
            date="2026-12-25",
            time="20:00",
            map_name="Bermuda",
            max_teams=16,
            registration_deadline=datetime.utcnow() + timedelta(days=7),
            rules="Fair play",
            status="OPEN",
        )
        db.session.add_all([admin, t])
        db.session.flush()

        r_conf = Registration(
            tournament_id=t.id,
            registration_id="PV-0001",
            team_name="Team Confirmed",
            captain_name="Cap 1",
            captain_phone="9876543211",
            payment_transaction_id="TXN1",
            payment_status="VERIFIED",
            registration_status="CONFIRMED",
            confirmation_date=datetime.utcnow(),
        )
        r_pend = Registration(
            tournament_id=t.id,
            team_name="Team Pending",
            captain_name="Cap 2",
            captain_phone="9876543212",
            payment_transaction_id="TXN2",
            payment_status="PAID_PENDING",
            registration_status="PENDING_PAYMENT_VERIFICATION",
        )
        r_rej = Registration(
            tournament_id=t.id,
            team_name="Team Rejected",
            captain_name="Cap 3",
            captain_phone="9876543213",
            payment_transaction_id="TXN3",
            payment_status="FAILED",
            registration_status="REJECTED",
            rejection_reason="Invalid payment slip",
        )
        db.session.add_all([r_conf, r_pend, r_rej])
        db.session.flush()

        # Add players for confirmed registration
        for i in range(1, 5):
            db.session.add(
                Player(
                    registration_id=r_conf.id,
                    player_index=i,
                    name=f"Player {i}",
                    ff_uid=f"100000000{i}",
                    ff_level=40 + i,
                )
            )

        db.session.commit()
        return t.id


def test_excel_export_unauthenticated(app, client):
    _setup_data(app)
    resp = client.get("/admin/export/excel", follow_redirects=False)
    assert resp.status_code in (301, 302, 401)


def test_excel_export_authenticated_and_filters(app, client):
    _setup_data(app)
    # Login admin
    client.post("/admin/login", data={"email": "admin@test.com", "password": "password123"})

    # 1. Export all
    resp_all = client.get("/admin/export/excel?filter=all")
    assert resp_all.status_code == 200
    wb_all = openpyxl.load_workbook(io.BytesIO(resp_all.data))
    ws_all = wb_all.active
    # Header + 3 registration rows = 4 rows
    assert ws_all.max_row == 4

    # Verify headers
    headers = [cell.value for cell in ws_all[1]]
    assert "Registration ID" in headers
    assert "Team Name" in headers
    assert "Captain Phone" in headers
    assert "Player 1 UID" in headers

    # 2. Export confirmed only
    resp_conf = client.get("/admin/export/excel?filter=confirmed")
    assert resp_conf.status_code == 200
    wb_conf = openpyxl.load_workbook(io.BytesIO(resp_conf.data))
    ws_conf = wb_conf.active
    # Header + 1 confirmed row = 2 rows
    assert ws_conf.max_row == 2
    row2 = [cell.value for cell in ws_conf[2]]
    assert "PV-0001" in row2
    assert "Team Confirmed" in row2

    # 3. Export pending only
    resp_pend = client.get("/admin/export/excel?filter=pending")
    assert resp_pend.status_code == 200
    wb_pend = openpyxl.load_workbook(io.BytesIO(resp_pend.data))
    ws_pend = wb_pend.active
    assert ws_pend.max_row == 2
    row2_pend = [cell.value for cell in ws_pend[2]]
    assert "Team Pending" in row2_pend

    # 4. Export rejected only
    resp_rej = client.get("/admin/export/excel?filter=rejected")
    assert resp_rej.status_code == 200
    wb_rej = openpyxl.load_workbook(io.BytesIO(resp_rej.data))
    ws_rej = wb_rej.active
    assert ws_rej.max_row == 2
    row2_rej = [cell.value for cell in ws_rej[2]]
    assert "Team Rejected" in row2_rej
    assert "Invalid payment slip" in row2_rej
